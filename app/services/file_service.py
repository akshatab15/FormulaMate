"""Upload validation and workbook analysis (xlsx / csv).

Uploaded files are only held in memory for analysis and are never stored
permanently. Macros and cell formulas are never executed.
"""

import csv
import io
from typing import List

from openpyxl import load_workbook

MAX_UPLOAD_MB = 5
ALLOWED_EXTENSIONS = (".xlsx", ".csv")
PREVIEW_ROWS = 8


class FileAnalysisError(Exception):
    """Raised for invalid, empty or unsupported uploads."""

    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.message = message
        self.status_code = status_code


def validate_upload(filename: str, size_bytes: int) -> str:
    name = (filename or "").strip()
    if not name:
        raise FileAnalysisError("No file was provided.", 400)
    lower = name.lower()
    if not lower.endswith(ALLOWED_EXTENSIONS):
        raise FileAnalysisError(
            "Unsupported file type. Please upload a .xlsx or .csv file.", 415
        )
    if size_bytes <= 0:
        raise FileAnalysisError("The uploaded file is empty.", 400)
    limit = MAX_UPLOAD_MB * 1024 * 1024
    if size_bytes > limit:
        raise FileAnalysisError(
            f"File is too large. Maximum allowed size is {MAX_UPLOAD_MB} MB.", 413
        )
    return "csv" if lower.endswith(".csv") else "xlsx"


def _clean(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _header_key(header: str) -> str:
    return re_sub_numbers(header.lower())


def re_sub_numbers(text: str) -> str:
    return "".join(ch for ch in text if ch.isalpha() or ch.isspace()).strip()


def _detect_types(rows: List[List[str]]) -> List[str]:
    """Guess a simple data type for every column."""
    if not rows:
        return []
    sample = rows[:20]
    types = []
    width = max(len(r) for r in sample)
    for col in range(width):
        values = [_cell(r, col) for r in sample]
        filled = [v for v in values if v != ""]
        if not filled:
            types.append("empty")
            continue
        numeric = sum(1 for v in filled if _is_number(v))
        date_like = sum(1 for v in filled if _looks_like_date(v))
        if date_like / len(filled) >= 0.6:
            types.append("date")
        elif numeric / len(filled) >= 0.6:
            types.append("number")
        else:
            types.append("text")
    return types


def _cell(row: List[str], index: int) -> str:
    return row[index] if index < len(row) else ""


def _is_number(value: str) -> bool:
    cleaned = value.replace(",", "").replace("%", "").strip()
    if cleaned.startswith("-"):
        cleaned = cleaned[1:]
    if cleaned.count(".") > 1:
        return False
    cleaned = cleaned.replace(".", "", 1)
    return cleaned.isdigit() and cleaned != ""


def _looks_like_date(value: str) -> bool:
    if "/" in value or "-" in value:
        parts = value.replace("-", "/").split("/")
        if len(parts) == 3 and all(p.strip().isdigit() for p in parts):
            return True
    return False


def _read_xlsx(data: bytes) -> List[dict]:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # corrupt or unreadable workbook
        raise FileAnalysisError(
            "The workbook could not be read. It may be corrupt or password protected.", 400
        ) from exc
    try:
        sheets = []
        for worksheet in workbook.worksheets:
            rows: List[List[str]] = []
            for row in worksheet.iter_rows(values_only=True):
                rows.append([_clean(cell) for cell in row])
                if len(rows) >= 60:
                    break
            sheets.append(_sheet_payload(worksheet.title, rows))
        return sheets
    finally:
        workbook.close()


def _read_csv(data: bytes) -> List[dict]:
    text = None
    for encoding in ("utf-8-sig", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise FileAnalysisError("The CSV file could not be decoded.", 400)
    if text.strip() == "":
        raise FileAnalysisError("The CSV file contains no data.", 400)

    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows = [ [_clean(cell) for cell in row] for row in reader ]
    if not rows or all(cell == "" for row in rows for cell in row):
        raise FileAnalysisError("The CSV file contains no data.", 400)
    return [_sheet_payload("CSV", rows)]


def _sheet_payload(title: str, rows: List[List[str]]) -> dict:
    rows = [row for row in rows if any(cell != "" for cell in row)]
    if not rows:
        return {
            "name": title,
            "headers": [],
            "types": [],
            "row_count": 0,
            "preview": [],
            "empty": True,
        }
    width = max(len(row) for row in rows[:20])
    normalised = [row[:width] + [""] * (width - len(row)) for row in rows]
    headers = [
        _clean(h) or f"Column {i + 1}" for i, h in enumerate(normalised[0])
    ]
    body = normalised[1:]
    types = _detect_types(body)
    return {
        "name": title,
        "headers": headers,
        "types": types,
        "row_count": len(body),
        "preview": [dict(zip(headers, row)) for row in body[:PREVIEW_ROWS]],
        "empty": False,
    }


def suggest_formulas(sheets: List[dict]) -> List[str]:
    """Map detected headers and data types to useful formula ids."""
    headers: List[str] = []
    types: List[str] = []
    for sheet in sheets:
        headers.extend(sheet["headers"])
        types.extend(sheet["types"])
    lowered = [h.lower() for h in headers]
    blob = " ".join(lowered)
    suggestions: List[str] = []

    def push(*ids: str) -> None:
        for fid in ids:
            if fid not in suggestions:
                suggestions.append(fid)

    if "number" in types:
        push("sum", "average")
    if any(word in blob for word in ("region", "category", "department", "team", "state")):
        push("sumif", "countif")
    if any(word in blob for word in ("region", "category", "department", "team")) and any(
        word in blob for word in ("status", "state", "approved", "result", "stage")
    ):
        push("sumifs", "countifs")
    if any(word in blob for word in ("id", "code", "sku", "roll", "number")):
        push("vlookup", "xlookup", "match")
    if any(word in blob for word in ("price", "rate", "amount", "salary", "cost", "total")):
        push("sumproduct", "round")
    if any(word in blob for word in ("qty", "quantity", "units", "stock", "opening", "closing")):
        push("product", "sumproduct")
    if "date" in types or any(word in blob for word in ("date", "dob", "joined", "start", "end")):
        push("today", "month", "datedif", "networkdays")
    if any(word in blob for word in ("name", "first", "last", "full", "product")):
        push("concat", "proper", "textjoin")
    if any(word in blob for word in ("status", "result", "grade", "pass", "remark")):
        push("if", "countif")
    if "text" in types:
        push("trim", "upper", "len")
    if "number" in types:
        push("max", "min")
    if len(suggestions) < 4:
        push("counta", "countblank", "if", "round")
    return suggestions[:8]


def analyze(filename: str, data: bytes) -> dict:
    kind = validate_upload(filename, len(data))
    sheets = _read_csv(data) if kind == "csv" else _read_xlsx(data)
    if not sheets:
        raise FileAnalysisError("The workbook contains no sheets.", 400)
    if all(sheet["empty"] for sheet in sheets):
        raise FileAnalysisError("The file has no data to preview.", 400)
    return {"kind": kind, "sheets": sheets}
